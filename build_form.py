from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('/Users/2_dogg/code/anqi')
DOCX = OUT / '关于同意申办出入境证件的函.docx'

BODY_FONT = 'FZFangSong-Z02'   # 仿宋
TITLE_FONT = 'Songti SC'       # 宋体
BODY_SIZE = 15.0
TITLE_SIZE = 25.0


def set_run_font(run, name=BODY_FONT, size=BODY_SIZE, bold=False, underline=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{attr}'), name)


def prep_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, left_indent_mm=0, before=0, after=0, line_mm=None):
    p.alignment = align
    pf = p.paragraph_format
    pf.left_indent = Mm(left_indent_mm) if left_indent_mm else None
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line_mm is not None:
        pf.line_spacing = Pt(line_mm * 2.83465)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    else:
        pf.line_spacing = 1
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def add_run(p, text, *, size=BODY_SIZE, font=BODY_FONT, bold=False, underline=False):
    r = p.add_run(text)
    set_run_font(r, font, size, bold, underline)
    return r


def add_spacer(doc, mm):
    p = doc.add_paragraph()
    prep_paragraph(p, line_mm=mm)
    add_run(p, ' ', size=1)
    return p


def add_fixed_line(doc, parts, height_mm=10, align=WD_ALIGN_PARAGRAPH.LEFT, indent_mm=0, tab_mm=None):
    p = doc.add_paragraph()
    prep_paragraph(p, align, indent_mm, line_mm=height_mm)
    if tab_mm is not None:
        p.paragraph_format.tab_stops.add_tab_stop(Mm(tab_mm), WD_TAB_ALIGNMENT.LEFT)
    for part in parts:
        if isinstance(part, str):
            add_run(p, part)
        else:
            text, opts = part
            add_run(p, text, **opts)
    return p


def set_table_fixed(table, width_mm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblLayout = tblPr.first_child_found_in('w:tblLayout')
    if tblLayout is None:
        tblLayout = OxmlElement('w:tblLayout')
        tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'), 'fixed')
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(width_mm * 56.6929)))
    tblW.set(qn('w:type'), 'dxa')


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'nil')


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_signature_paragraph(doc, left, right, height_mm=11):
    p = doc.add_paragraph()
    prep_paragraph(p, line_mm=height_mm)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(124), WD_TAB_ALIGNMENT.LEFT)
    add_run(p, left)
    add_run(p, '\t')
    add_run(p, right)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(43)
    sec.bottom_margin = Mm(14)
    sec.left_margin = Mm(23)
    sec.right_margin = Mm(21)
    sec.header_distance = Mm(5)
    sec.footer_distance = Mm(5)

    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal._element.rPr.rFonts.set(qn('w:ascii'), BODY_FONT)
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), BODY_FONT)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    normal._element.rPr.rFonts.set(qn('w:cs'), BODY_FONT)

    # 标题
    p = doc.add_paragraph()
    prep_paragraph(p, left_indent_mm=29, line_mm=12)
    add_run(p, '关于同意', size=TITLE_SIZE, font=TITLE_FONT, bold=True)
    add_run(p, '　' * 16, size=TITLE_SIZE, font=TITLE_FONT, bold=True, underline=True)
    p = doc.add_paragraph()
    prep_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, line_mm=12)
    add_run(p, '申办出入境证件的函', size=TITLE_SIZE, font=TITLE_FONT, bold=True)

    add_spacer(doc, 13)

    # 正文及填空线（下划线空格可以直接覆盖输入）
    add_fixed_line(doc, [('　' * 11, {'underline': True}), '出入境管理部门：'], 10)
    add_fixed_line(doc, [
        '　' * 3,
        ('同志（身份证号码：', {}),
        ('　' * 17, {'underline': True}),
        '）',
    ], 10)
    add_fixed_line(doc, ['系', ('　' * 20, {'underline': True}), '（填写单位全称）'], 10)
    add_fixed_line(doc, ['的', ('　' * 13, {'underline': True}), '（填写职务）。按照干部管理权限，'], 10)

    # 勾选项
    add_fixed_line(doc, ['我单位同意该人申办：  □普通护照', '\t', '□往来港澳通行证及签注'], 10, tab_mm=104)
    add_fixed_line(doc, ['□往来台湾通行证及签注。'], 10)

    add_spacer(doc, 17)

    # 联系人
    add_fixed_line(doc, ['组织、人事部门联系人姓名：'], 10, align=WD_ALIGN_PARAGRAPH.LEFT, indent_mm=15)
    add_fixed_line(doc, ['联系电话：'], 10, align=WD_ALIGN_PARAGRAPH.CENTER, indent_mm=8)

    add_spacer(doc, 16)

    # 负责人、盖章、日期
    add_signature_paragraph(doc, '负责人签名：', '公    章', 11)
    add_signature_paragraph(doc, '', '年    月    日', 11)

    add_spacer(doc, 19)

    # 备注
    add_fixed_line(doc, ['备注：1、登记备案国家工作人员申请出入境证件须提交此函。'], 9)
    add_fixed_line(doc, ['2、登记备案单位须在同意办理的出入境证件类型前打钩，'], 9, indent_mm=9)
    add_fixed_line(doc, ['并划掉不同意办理的证件类型。'], 9, indent_mm=24)
    add_fixed_line(doc, ['3、本函自开具之日起3个月内有效。'], 9, indent_mm=9)

    doc.core_properties.title = '关于同意申办出入境证件的函'
    doc.core_properties.subject = '出入境证件申请同意函空白表格'
    doc.core_properties.author = ''
    doc.core_properties.comments = ''
    doc.save(DOCX)
    print(DOCX)

if __name__ == '__main__':
    main()
